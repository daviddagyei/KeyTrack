import openpyxl

# Load the workbook
wb = openpyxl.load_workbook('key_distribution.xlsx')
sheet = wb.active

# Print room IDs
print("Room IDs in Excel file:")
for row in range(2, min(12, sheet.max_row + 1)):
    room_id = sheet.cell(row=row, column=1).value
    print(f'Row {row}: "{room_id}"')
