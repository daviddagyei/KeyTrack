import openpyxl

# Load the workbook
wb = openpyxl.load_workbook('key_distribution.xlsx')
sheet = wb.active

# Print headers
print("Headers:")
for col in range(1, 7):
    print(f"Column {col}: {sheet.cell(row=1, column=col).value}")

# Print sample data
print("\nSample Data:")
for row in range(2, min(5, sheet.max_row + 1)):
    room_id = sheet.cell(row=row, column=1).value
    total_keys = sheet.cell(row=row, column=2).value
    collected = sheet.cell(row=row, column=3).value or ""
    lost = sheet.cell(row=row, column=4).value or ""
    borrowed = sheet.cell(row=row, column=5).value or ""
    returned = sheet.cell(row=row, column=6).value or ""
    
    print(f"Room: {room_id}")
    print(f"  - Total Keys: {total_keys}")
    print(f"  - Collected: {collected}")
    print(f"  - Lost: {lost}")
    print(f"  - Borrowed: {borrowed}")
    print(f"  - Returned: {returned}")
    print()

# Print total rooms
print(f"Total rooms in spreadsheet: {sheet.max_row - 1}")
