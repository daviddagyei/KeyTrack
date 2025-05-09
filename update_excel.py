import openpyxl
from datetime import datetime


print("Loading workbook...")
wb = openpyxl.load_workbook('key_distribution.xlsx')
sheet = wb.active

print("Updating room data according to new logic...")
for row in range(2, sheet.max_row + 1):
    room_id = sheet.cell(row=row, column=1).value
    if not room_id:
        continue

    collected_str = sheet.cell(row=row, column=3).value or ""
    lost_str = sheet.cell(row=row, column=4).value or ""
    borrowed_str = sheet.cell(row=row, column=5).value or ""
    returned_str = sheet.cell(row=row, column=6).value or ""
    

    collected_count = len([x for x in collected_str.split(',') if x.strip()])
    lost_count = len([x for x in lost_str.split(',') if x.strip()])
    borrowed_count = len([x for x in borrowed_str.split(',') if x.strip()])
    returned_count = len([x for x in returned_str.split(',') if x.strip()])

    total_keys = collected_count + lost_count
    
    active_collected = collected_count - min(returned_count, collected_count)
    
    available_keys = total_keys - (active_collected + borrowed_count)
    
    if available_keys < 0:
        total_keys += abs(available_keys)
        available_keys = 0
        
    print(f"Room {room_id}: Total={total_keys}, Available={available_keys}, Active Collected={active_collected}, Borrowed={borrowed_count}")
    
    sheet.cell(row=row, column=2, value=total_keys)

wb.save('key_distribution_updated.xlsx')
print("Workbook saved as key_distribution_updated.xlsx")
print("Please verify the updates and replace the original file if satisfied.")