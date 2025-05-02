import openpyxl

# Load the workbook
wb = openpyxl.load_workbook('key_distribution.xlsx')
sheet = wb.active

# Print the first 10 room IDs from the Excel file
print("Room IDs in Excel file:")
for row in range(2, min(12, sheet.max_row + 1)):
    room_id = sheet.cell(row=row, column=1).value
    print(f"Row {row}: '{room_id}'")

